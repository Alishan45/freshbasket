import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# ==============================================================================
# AUDITING SUITE FOR STRICT WRITING RULES
# ==============================================================================
BANNED_WORDS = [
    'delve', 'tapestry', 'testament', 'pivotal', 'underscore', 'intricate',
    'meticulously', 'vibrant', 'realm', 'beacon', 'multifaceted', 'foster',
    'showcase', 'bolster', 'garner', 'interplay', 'enduring', 'crucial',
    'vital', 'landscape', 'dynamic', 'nuanced', 'watershed moment',
    'turning point', 'deeply rooted', 'rich heritage', 'steadfast',
    'paradigm', 'synergy', 'leverage'
]

BANNED_TRANSITIONS = [
    'moreover', 'furthermore', 'in addition', 'consequently', 'nonetheless',
    'it is important to note', 'it is worth mentioning', 'notably', 'crucially',
    "in today's fast-paced world"
]

BANNED_VERBS = [
    'serves as', 'stands as', 'represents', 'marks', 'boasts'
]

BANNED_PHRASES = [
    'not only', 'not just', 'rather than', 'groundbreaking', 'breathtaking',
    'remarkable', 'seamless', 'transformative', 'studies show', 'experts say',
    'researchers argue', 'observers note'
]

def audit_text(text, context=""):
    errors = []
    lower = text.lower()
    for w in BANNED_WORDS:
        if re.search(r'\b' + re.escape(w) + r'\b', lower):
            errors.append(f"[{context}] Banned word: '{w}'")
    for t in BANNED_TRANSITIONS:
        if re.search(r'\b' + re.escape(t) + r'\b', lower):
            errors.append(f"[{context}] Banned transition: '{t}'")
    for v in BANNED_VERBS:
        if re.search(r'\b' + re.escape(v) + r'\b', lower):
            errors.append(f"[{context}] Banned verb: '{v}'")
    for p in BANNED_PHRASES:
        if re.search(r'\b' + re.escape(p) + r'\b', lower):
            errors.append(f"[{context}] Banned phrase: '{p}'")
    if '—' in text:
        errors.append(f"[{context}] Em dash found")
    if re.search(r'\*\*[^*]+:\*\*', text):
        errors.append(f"[{context}] Bold Term explanation format found")
    return errors

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def build_task2_document():
    doc = Document()
    all_audit_errors = []

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(5)

    def add_p(text, bold=False, italic=False, color=None, space_after=6):
        errs = audit_text(text, "paragraph")
        if errs:
            all_audit_errors.extend(errs)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(10)
        run.font.bold = bold
        run.font.italic = italic
        if color:
            run.font.color.rgb = RGBColor(*color)
        else:
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        return p

    def add_heading1(text):
        errs = audit_text(text, "heading 1")
        if errs:
            all_audit_errors.extend(errs)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x14, 0x53, 0x2D)
        return p

    def add_heading2(text):
        errs = audit_text(text, "heading 2")
        if errs:
            all_audit_errors.extend(errs)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        return p

    def add_heading3(text):
        errs = audit_text(text, "heading 3")
        if errs:
            all_audit_errors.extend(errs)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x04, 0x78, 0x57)
        return p

    def format_table(table, col_widths, headers, rows_data):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="E2E8F0")

        hdr_row = table.rows[0]
        hdr_tr = hdr_row._tr.get_or_add_trPr()
        hdr_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for idx, text in enumerate(headers):
            cell = hdr_row.cells[idx]
            cell.width = col_widths[idx]
            set_cell_shading(cell, "14532D")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        for r_idx, row_data in enumerate(rows_data):
            row = table.rows[r_idx + 1]
            bg_hex = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
            for c_idx, val in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.width = col_widths[c_idx]
                set_cell_shading(cell, bg_hex)
                set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(str(val))
                run.font.name = 'Arial'
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    def add_image_box(img_path, caption, width_in=6.2):
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            run = p_img.add_run()
            run.add_picture(img_path, width=Inches(width_in))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(10)
            r_cap = p_cap.add_run(f"Figure: {caption}")
            r_cap.font.name = 'Arial'
            r_cap.font.size = Pt(8.5)
            r_cap.font.italic = True
            r_cap.font.bold = True
            r_cap.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        else:
            add_p(f"[Placeholder: {caption} - File {img_path} not found]")

    # ==========================================================================
    # COVER / TITLE PAGE (Matches task2_original_template.docx placeholders)
    # ==========================================================================
    for _ in range(2):
        doc.add_paragraph()

    p_t1 = doc.add_paragraph()
    p_t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t1.paragraph_format.space_after = Pt(4)
    r1 = p_t1.add_run("FreshBasket")
    r1.font.name = 'Arial'
    r1.font.size = Pt(30)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x14, 0x53, 0x2D)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    r_sub = p_sub.add_run("Online fresh fruits, vegetables, and essential grocery delivery platform")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(0x15, 0x80, 0x3D)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("____________________________________________________")
    r_div.font.color.rgb = RGBColor(0x86, 0xEF, 0xAC)
    p_div.paragraph_format.space_after = Pt(24)

    p_deliv = doc.add_paragraph()
    p_deliv.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_deliv.paragraph_format.space_after = Pt(6)
    r_deliv = p_deliv.add_run("Documentation (first deliverable)")
    r_deliv.font.name = 'Arial'
    r_deliv.font.size = Pt(18)
    r_deliv.font.bold = True
    r_deliv.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_srs = doc.add_paragraph()
    p_srs.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_srs.paragraph_format.space_after = Pt(28)
    r_srs = p_srs.add_run("Software requirements specification and system design")
    r_srs.font.name = 'Arial'
    r_srs.font.size = Pt(12)
    r_srs.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    meta_tbl = doc.add_table(rows=6, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Document version:", "Version 1.0"),
        ("Course title:", "CS619 Final Project"),
        ("Group ID:", "F26PROJECT56E94"),
        ("Student ID:", "BC250206138"),
        ("Supervisor name:", "Muhammad Hassaan"),
        ("Supervisor email:", "m.hassaan@vu.edu.pk")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_tbl.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(3.8)
        c0.text = label
        c1.text = val
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        c1.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x14, 0x53, 0x2D)
        set_cell_margins(c0, 50, 50, 80, 80)
        set_cell_margins(c1, 50, 50, 80, 80)

    for _ in range(3):
        doc.add_paragraph()

    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = p_foot.add_run("Department of Computer Science\nVirtual University of Pakistan\nSemester Fall 2026")
    r_foot.font.size = Pt(9.5)
    r_foot.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    doc.add_page_break()

    # ==========================================================================
    # REVISION HISTORY & TABLE OF CONTENTS
    # ==========================================================================
    add_heading1("Revision history")
    add_p("The table below records the revision versions and development notes for this first deliverable specification.")

    rev_tbl = doc.add_table(rows=4, cols=4)
    rev_headers = ["Date (dd/mm/yyyy)", "Version", "Description", "Author"]
    rev_data = [
        ["28/10/2026", "1.0", "Initial complete first deliverable: SRS, use cases, work plan Gantt, ERD, database schema, and interface designs", "BC250206138"],
        ["09/11/2026", "1.1", "Formal submission of Documentation deliverable to Virtual University portal", "BC250206138"],
        ["07/01/2027", "2.0", "Final deliverable source code, database scripts, and project defense submission", "BC250206138"]
    ]
    format_table(rev_tbl, [Inches(1.5), Inches(0.8), Inches(2.7), Inches(1.5)], rev_headers, rev_data)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    add_heading1("Table of contents")
    toc_items = [
        ("1. Introduction of document", "3"),
        ("   1.1 Purpose of the document", "3"),
        ("   1.2 Purpose and benefits of the first deliverable phase", "3"),
        ("   1.3 Structural inclusions of this specification", "4"),
        ("2. Scope (of the project)", "5"),
        ("   2.1 Problem definition and background", "5"),
        ("   2.2 System goals", "5"),
        ("   2.3 User roles and responsibilities", "6"),
        ("   2.4 Functional boundaries and limits", "6"),
        ("   2.5 Technology selection", "7"),
        ("3. Functional requirements and non-functional requirements", "8"),
        ("   3.1 Detailed functional requirements (FR-01 to FR-10)", "8"),
        ("   3.2 Comprehensive non-functional requirements (NFR-01 to NFR-05)", "11"),
        ("4. Use case diagram", "13"),
        ("   4.1 System actors and operational roles", "13"),
        ("   4.2 Use case diagram", "14"),
        ("   4.3 Use case tabular specifications", "15"),
        ("5. Work plan", "18"),
        ("   5.1 Work breakdown structure and calendar phases", "18"),
        ("   5.2 Project Gantt chart", "19"),
        ("   5.3 Milestone schedule and deliverables table", "20"),
        ("6. Entity relationship diagram (ERD)", "21"),
        ("   6.1 Conceptual data model and cardinalities", "21"),
        ("   6.2 Entity relationship diagram", "22"),
        ("   6.3 Entity dictionary and attribute definitions", "23"),
        ("7. Database design", "25"),
        ("   7.1 Relational database schema structure", "25"),
        ("   7.2 Physical relational schema diagram", "26"),
        ("   7.3 Complete database table specifications", "27"),
        ("8. Interface design", "31"),
        ("   8.1 User interface layout and responsive principles", "31"),
        ("   8.2 Storefront catalog and hub selector", "32"),
        ("   8.3 Produce details and farm origin display", "33"),
        ("   8.4 Shopping cart and cash on delivery checkout", "34"),
        ("   8.5 Real-time five-stage order tracking portal", "35"),
        ("   8.6 Administrative dashboard and logistics console", "36"),
    ]

    for item, page in toc_items:
        p_toc = doc.add_paragraph()
        p_toc.paragraph_format.space_after = Pt(2)
        p_toc.paragraph_format.line_spacing = 1.15
        is_main = not item.startswith("   ")
        r_item = p_toc.add_run(item)
        r_item.font.bold = is_main
        r_item.font.size = Pt(9.5 if is_main else 9.0)
        r_item.font.color.rgb = RGBColor(0x14, 0x53, 0x2D) if is_main else RGBColor(0x33, 0x41, 0x55)

        dots_count = max(4, 76 - len(item))
        r_dots = p_toc.add_run(" " + "." * dots_count + " ")
        r_dots.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)

        r_page = p_toc.add_run(page)
        r_page.font.bold = is_main
        r_page.font.size = Pt(9.0)
        r_page.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_page_break()

    # ==========================================================================
    # SECTION 1: INTRODUCTION OF DOCUMENT
    # ==========================================================================
    add_heading1("1. Introduction of document")

    add_heading2("1.1 Purpose of the document")
    add_p("This document is the official first deliverable for the FreshBasket project. It specifies the functional requirements, quality criteria, structural diagrams, and interface layouts for the web application. The specification forms a technical agreement between student BC250206138, supervisor Muhammad Hassaan, and the Virtual University evaluation committee. System behavior is defined before implementation starts. The text translates general project goals into exact software modules, data entities, process workflows, and user screens.")

    add_heading2("1.2 Purpose and benefits of the first deliverable phase")
    add_p("The requirements and design phase establishes the foundation of the development lifecycle. In the 1981 textbook Software Engineering Economics, Barry Boehm documented that finding and fixing software errors during early requirements analysis costs 50 to 100 times less than correcting defects after deployment. Defining system scope and data structures before programming prevents architectural rework.")
    add_p("The first deliverable phase provides five concrete benefits. First, it defines clear project boundaries. The document specifies what functions belong in the system and what features are excluded. Second, it aligns the development plan with project supervisor Muhammad Hassaan. The supervisor evaluates technical rigor and feasibility before coding begins. Third, it establishes architectural traceability. Every interface screen and database table connects directly to a numbered requirement. Fourth, it provides an implementation blueprint. Database tables, column types, and foreign key rules give the developer exact instructions. Fifth, it supplies measurable quality baselines. Concrete performance, security, and usability criteria support final project evaluation.")

    add_heading2("1.3 Structural inclusions of this specification")
    add_p("This document contains eight cohesive sections. The scope section explains agricultural distribution problems, user roles, system limits, and technology components. The requirements section details ten functional modules and five quality attributes. The use case section defines four operational roles and provides tabular scenario specifications. The work plan presents an 11-week academic schedule with milestone deadlines. The entity relationship section details 14 normalized entities and data cardinalities. The database design section provides physical table specifications and foreign key rules. The interface design section contains screenshots and walkthroughs of five primary application screens.")

    doc.add_page_break()

    # ==========================================================================
    # SECTION 2: SCOPE (OF THE PROJECT)
    # ==========================================================================
    add_heading1("2. Scope (of the project)")

    add_heading2("2.1 Problem definition and background")
    add_p("Fresh produce distribution in Pakistan faces physical and organizational bottlenecks. Fruits and vegetables pass through multiple wholesale commission agents, open-air mandis, and ambient roadside vendors over three to six days. This extended transit time causes up to 40 percent post-harvest spoilage. Consumers encounter arbitrary daily price markups, inaccurate weights, and artificial chemical ripening agents like calcium carbide.")
    add_p("FreshBasket addresses these problems through a direct online grocery marketplace. The platform connects urban households with fresh produce sourced directly from regional farming clusters and controlled cold-chain centers. Orders placed on the website are packed at regional hubs and dispatched to customer doorsteps in temperature-monitored crates within 24 hours of harvest.")

    add_heading2("2.2 System goals")
    add_p("The FreshBasket application satisfies four core objectives. First, the platform enables direct farm sourcing. Urban consumers receive seasonal fruits and vegetables harvested within 24 hours. Second, the system supports real-world grocery purchasing by weight in kilograms and grams, count in individual pieces, or volume in dozens with accurate price calculation. Third, the application provides order transparency. Deliveries progress through five verifiable stages with courier location updates. Fourth, the system supplies operational tools for store managers. Administrators and logistics supervisors inspect produce quality, manage inventory stock, monitor sales, and resolve customer complaints.")

    add_heading2("2.3 User roles and responsibilities")
    add_p("The platform serves four user roles with distinct operational permissions. Customers search the produce catalog, filter items by category, detect location via Geoapify geolocation, build shopping carts, place Cash on Delivery orders, track deliveries, write reviews, and file complaints. Store administrators manage products and categories, set inventory thresholds, inspect produce batches, moderate customer reviews, and answer complaint tickets. Logistics supervisors monitor active orders, assign delivery couriers, advance fulfillment stages, and coordinate regional dispatch hubs. Farmers and suppliers register harvest lots, update produce availability, and track batch receipts at local collection points.")

    add_heading2("2.4 Functional boundaries and limits")
    add_p("The functional scope includes user account registration, profile management, and multiple saved delivery addresses. It includes an online catalog of fruits and vegetables with nutritional details, freshness grades, and harvest dates. It provides a multi-unit shopping cart with automatic price recalculation and discount validation. It supports Cash on Delivery checkout with delivery time-slot selection and special delivery instructions. It offers real-time five-stage fulfillment tracking with courier route visualization. It includes Geoapify geolocation services for automatic address resolution and regional hub detection. It provides customer complaint ticket management with automated Gmail SMTP email dispatch and post-delivery product ratings.")
    add_p("The project excludes third-party electronic payment gateway integration. Cash on Delivery is the sole payment method. The project excludes cross-border international delivery and automated drone delivery hardware.")

    add_heading2("2.5 Technology selection")
    add_p("The FreshBasket project uses modern web engineering tools across the full application stack. The frontend uses React 18 with the Vite build tool, React Router 6, and Tailwind CSS. The backend uses Node.js with the Express web framework and modular route controllers. The database uses SQLite with the node:sqlite driver, foreign key constraints, and 14 normalized tables. Security uses Argon2 and bcrypt password hashing with JSON Web Tokens for role-based route protection. External integrations use Geoapify for reverse geocoding and static map generation. Nodemailer handles complaint notification emails via Gmail SMTP.")

    doc.add_page_break()

    # ==========================================================================
    # SECTION 3: FUNCTIONAL & NON-FUNCTIONAL REQUIREMENTS
    # ==========================================================================
    add_heading1("3. Functional requirements and non-functional requirements")

    add_heading2("3.1 Detailed functional requirements")

    add_heading3("FR-01: User registration and account management")
    add_p("The system allows new users to register using full name, email address, password, and phone number. The backend hashes all passwords using Argon2 or bcrypt before storing them in the database. Registered users log in, maintain personal profile data, save multiple delivery addresses with GPS coordinates, and view previous purchase history.")

    add_heading3("FR-02: Fruit and vegetable catalog management")
    add_p("The platform provides a catalog of fresh produce organized into categories: Fruits, Vegetables, Leafy Greens, Organic, and Seasonal. Store administrators create, update, and deactivate products. Metadata fields include title, description, category, base price, inventory count, freshness grade, freshness score, farm origin, and harvest date.")

    add_heading3("FR-03: Measurement units and unit price calculation")
    add_p("The system allows customers to select products using real-world grocery measurement units. Products support kilograms, grams (250g, 500g), individual pieces, or dozens. Each unit configuration contains a numerical price multiplier. The backend recalculates item subtotals from database records to prevent client-side price tampering.")

    add_heading3("FR-04: Shopping cart and cash on delivery checkout")
    add_p("Customers add produce to a shopping cart, modify item quantities, and remove products. The checkout screen presents an itemized billing summary with subtotal, delivery fee, and applicable discounts. Customers select a saved delivery address or enter a new doorstep location, choose a delivery time window, add special instructions, and confirm their order using Cash on Delivery.")

    add_heading3("FR-05: Real-time five-stage order fulfillment")
    add_p("Every placed order progresses through five verified fulfillment stages. In the pending stage, the customer places the order and the database logs the record. In the confirmed stage, store staff validate stock availability and customer address. In the packed stage, staff select produce from temperature-controlled bins and pack items into crates. In the out for delivery stage, the courier collects the package and proceeds along the transit route. In the delivered stage, the courier delivers the produce to the doorstep and collects cash payment. Every status change writes a record to the order_status_history table with a timestamp and the user ID of the updater.")

    add_heading3("FR-06: Logistics supervisor dispatch portal")
    add_p("A dedicated portal allows logistics supervisors to oversee order fulfillment. Supervisors view active orders filtered by status, assign delivery couriers, record dispatch notes, and advance order stages. Updates broadcast in real time across customer tracking screens, administrator panels, and supervisor dashboards.")

    add_heading3("FR-07: Geoapify geolocation and regional hub selection")
    add_p("The platform integrates Geoapify API services using key 884ee10ba0034a499dacd7a6359821ce. The frontend uses browser coordinates to perform reverse geocoding, returning street, area, and city names. The Haversine distance formula determines the nearest regional cold-chain hub from five facilities: Lahore Central, DHA Cantt, Johar Town, Islamabad Margalla, and Karachi Clifton. The checkout and tracking screens display static Geoapify map images showing customer doorstep pins and transit routes.")

    add_heading3("FR-08: Customer reviews and ratings")
    add_p("Customers who receive a completed order submit star ratings (1 to 5 stars) and written reviews for individual products. The backend verifies order completion before accepting review submissions. Store administrators moderate reviews and remove inappropriate entries.")

    add_heading3("FR-09: Email complaint center and dispute resolution")
    add_p("The platform provides a customer complaint center. Customers submit dispute tickets with category selection (Quality, Missing Item, Late Delivery, Billing), order references, and descriptions. The system generates a unique ticket number (FB-TKT-XXXXX) and dispatches an automated HTML receipt email to the customer using Gmail SMTP (muzammilshabbircs01@gmail.com). Administrators view tickets on a dispute dashboard, compose replies, and send resolution emails directly to customers.")

    add_heading3("FR-10: Administrative governance and inventory auditing")
    add_p("The admin dashboard presents store metrics including gross revenue, order volume, active customer counts, and low-stock produce warnings. Every stock adjustment logs a record to the inventory_logs table containing the actor ID, quantity changed, reason, and resulting balance.")

    add_heading2("3.2 Comprehensive non-functional requirements")

    add_heading3("NFR-01: Performance and responsiveness")
    add_p("The storefront catalog and product search pages load within 2.0 seconds over standard broadband connections. Backend REST API endpoints respond within 200 milliseconds under a concurrent load of 50 active users. Database queries execute within 50 milliseconds by indexing foreign keys and search columns.")

    add_heading3("NFR-02: Information security and access control")
    add_p("All passwords are encrypted using Argon2 or bcrypt algorithms with unique salts. Protected API routes verify JSON Web Tokens signed with secret keys. The backend enforces role-based access control, preventing customers from accessing administrative or supervisor endpoints. Database queries use parameterized statements to eliminate SQL injection vulnerabilities.")

    add_heading3("NFR-03: Reliability and transactional integrity")
    add_p("Order placement and inventory deduction run inside database transactions. If an item stock check fails, the transaction rolls back completely to prevent overselling. The system handles external API network timeouts gracefully by falling back to stored local defaults.")

    add_heading3("NFR-04: Usability and mobile responsiveness")
    add_p("The user interface uses Tailwind CSS responsive breakpoints (sm, md, lg, xl). The layout adjusts cleanly across mobile phone screens (375px to 640px), tablet displays (768px to 1024px), and desktop monitors (1280px and wider). Navigation touch targets meet the 44px minimum sizing standard for touchscreens.")

    add_heading3("NFR-05: Maintainability and modularity")
    add_p("The codebase separates concerns between client-side user interface components and server-side API controllers. Frontend components reside in modular folders (/components, /pages, /services). Backend endpoints follow REST conventions and include automated unit test suites using the Node test runner.")

    doc.add_page_break()

    # ==========================================================================
    # SECTION 4: USE CASE DIAGRAM
    # ==========================================================================
    add_heading1("4. Use case diagram")

    add_heading2("4.1 System actors and operational roles")
    add_p("The system contains four distinct human actors. Customers browse produce, manage cart items, place Cash on Delivery orders, track deliveries, rate items, and file complaints. Administrators oversee product data, review sales metrics, audit inventory levels, and resolve complaint tickets. Logistics supervisors assign couriers, update five-stage fulfillment milestones, and coordinate regional hubs. Farmers and suppliers register harvest lots and update incoming produce stock.")

    add_heading2("4.2 Use case diagram")
    add_p("The diagram below illustrates the system boundary, actors, core use cases, and relationship dependencies.")
    add_image_box("diagrams/use_case_diagram.png", "System Use Case Diagram with Four Actors and Core Interactions", width_in=6.2)

    add_heading2("4.3 Use case tabular specifications")

    add_heading3("Use case UC-01: Place grocery order via cash on delivery")
    uc1_tbl = doc.add_table(rows=8, cols=2)
    uc1_data = [
        ("Use Case ID & Title:", "UC-01: Place Grocery Order via Cash on Delivery"),
        ("Primary Actor:", "Customer"),
        ("Preconditions:", "Customer is authenticated and has at least one produce item in the shopping cart."),
        ("Postconditions:", "Order is created in database with 'pending' status, inventory is deducted, and confirmation email is sent."),
        ("Main Success Flow:", "1. Customer opens shopping cart and reviews item quantities and measurement units.\n2. Customer clicks 'Proceed to Checkout'.\n3. System retrieves customer delivery addresses and detects nearest cold-chain hub via Geoapify.\n4. Customer selects delivery address, chooses delivery time slot, and confirms COD payment.\n5. Customer clicks 'Place Order'.\n6. Backend verifies product stock in a database transaction.\n7. Backend inserts record into orders table and order_items table.\n8. Backend deducts item stock from products table and creates inventory log.\n9. System returns order confirmation number and redirects customer to tracking page."),
        ("Alternative Flow A:", "Insufficient Stock: If an item stock falls below cart quantity, system aborts transaction and displays available quantity warning."),
        ("Alternative Flow B:", "Invalid Address: If customer enters an incomplete address, system highlights missing fields and prompts correction.")
    ]
    format_table(uc1_tbl, [Inches(1.8), Inches(4.4)], ["Attribute", "Specification Details"], uc1_data)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_heading3("Use case UC-02: Advance five-stage order fulfillment")
    uc2_tbl = doc.add_table(rows=7, cols=2)
    uc2_data = [
        ("Use Case ID & Title:", "UC-02: Advance Five-Stage Order Fulfillment"),
        ("Primary Actor:", "Logistics Supervisor / Administrator"),
        ("Preconditions:", "Supervisor is authenticated and order exists with status other than 'delivered' or 'cancelled'."),
        ("Postconditions:", "Order status is updated, status history record is inserted, and customer tracking view updates."),
        ("Main Success Flow:", "1. Supervisor opens Logistics Dispatch Portal.\n2. System loads active orders with current stages and destination hubs.\n3. Supervisor selects an order and chooses the next stage (e.g., 'packed' or 'out_for_delivery').\n4. Supervisor optionally enters courier name, contact number, and dispatch notes.\n5. Supervisor submits update.\n6. Backend verifies supervisor permissions and updates order_status column.\n7. Backend records change in order_status_history with timestamp.\n8. Real-time event bus notifies connected customer tracking screens."),
        ("Alternative Flow:", "Order Cancellation: If customer requests cancellation before packing, supervisor sets status to 'cancelled' and inventory is restored.")
    ]
    format_table(uc2_tbl, [Inches(1.8), Inches(4.4)], ["Attribute", "Specification Details"], uc2_data)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_heading3("Use case UC-03: Submit customer complaint ticket")
    uc3_tbl = doc.add_table(rows=7, cols=2)
    uc3_data = [
        ("Use Case ID & Title:", "UC-03: Submit Customer Complaint Ticket"),
        ("Primary Actor:", "Customer"),
        ("Preconditions:", "Customer is logged in and has an active or completed order reference."),
        ("Postconditions:", "Complaint ticket is created, automated receipt email is dispatched, and ticket appears on admin dashboard."),
        ("Main Success Flow:", "1. Customer navigates to Complaint Center (/complaint).\n2. System loads customer contact information and recent orders.\n3. Customer selects dispute category (Quality, Missing Item, Delivery Delay, Billing).\n4. Customer inputs subject, description, and optional photo attachment URL.\n5. Customer submits complaint.\n6. Backend generates ticket number (FB-TKT-XXXXX) and inserts complaint record.\n7. Backend triggers Nodemailer to send confirmation email via Gmail SMTP.\n8. System displays ticket tracking card with current status 'open'."),
        ("Alternative Flow:", "Missing Required Fields: System validates form and highlights empty subject or description fields.")
    ]
    format_table(uc3_tbl, [Inches(1.8), Inches(4.4)], ["Attribute", "Specification Details"], uc3_data)

    doc.add_page_break()

    # ==========================================================================
    # SECTION 5: WORK PLAN
    # ==========================================================================
    add_heading1("5. Work plan")

    add_heading2("5.1 Work breakdown structure and calendar phases")
    add_p("The project schedule spans 11 academic calendar weeks from Monday 26 October 2026 to Thursday 07 January 2027. The schedule corresponds to the two official Virtual University course milestones. Assignment 1 focuses on documentation from 26 October 2026 to 09 November 2026. This period covers problem analysis, requirements specification, use case modeling, ERD design, database normalization, and interface wireframing. Assignment 2 covers the final deliverable from 10 November 2026 to 07 January 2027. This phase includes full-stack implementation, REST APIs, database integration, geolocation services, logistics dispatching, testing, and the viva presentation.")

    add_heading2("5.2 Project Gantt chart")
    add_p("The Gantt chart below displays the chronological schedule, task durations, and milestones across the 11-week project lifecycle.")
    add_image_box("diagrams/work_plan_gantt.png", "Work Plan and Gantt Chart for Assignment 1 and Assignment 2", width_in=6.2)

    add_heading2("5.3 Milestone schedule and deliverables table")
    add_p("The table below details the project tasks, start dates, completion deadlines, dependencies, and expected deliverables.")

    wp_tbl = doc.add_table(rows=12, cols=5)
    wp_headers = ["Phase", "Task Description", "Start Date", "End Date", "Milestone Deliverable"]
    wp_data = [
        ["1.1", "Literature Review & Problem Formulation", "26/10/2026", "01/11/2026", "Scope Definition Document"],
        ["1.2", "SRS & Functional Requirements", "02/11/2026", "04/11/2026", "Draft SRS Document"],
        ["1.3", "Use Case Modeling & System Architecture", "04/11/2026", "06/11/2026", "Use Case Diagram & Scenarios"],
        ["1.4", "ERD & Database Schema Design", "06/11/2026", "08/11/2026", "Normalized 14-Table Schema"],
        ["1.5", "UI Wireframing & Documentation Finalization", "08/11/2026", "09/11/2026", "Assignment 1 Submission (09 Nov)"],
        ["2.1", "Full-Stack Setup & Authentication Engine", "10/11/2026", "19/11/2026", "React + Express Skeleton & Auth"],
        ["2.2", "Produce Catalog & Multi-Unit Pricing Engine", "20/11/2026", "29/11/2026", "Catalog Filters & Multi-Units"],
        ["2.3", "Cart, COD Checkout & Geoapify Geolocation", "30/11/2026", "09/12/2026", "Doorstep Pin & Checkout Flow"],
        ["2.4", "Real-Time 5-Stage Logistics & Dispatch Portal", "10/12/2026", "18/12/2026", "Supervisor Portal & Live Tracker"],
        ["2.5", "Customer Reviews & SMTP Complaint Center", "19/12/2026", "26/12/2026", "Review Engine & Email Tickets"],
        ["2.6", "Mobile Responsive QA, Final Viva & Defense", "27/12/2026", "07/01/2027", "Assignment 2 Submission (07 Jan)"]
    ]
    format_table(wp_tbl, [Inches(0.6), Inches(2.2), Inches(1.0), Inches(1.0), Inches(1.4)], wp_headers, wp_data)

    doc.add_page_break()

    # ==========================================================================
    # SECTION 6: ENTITY RELATIONSHIP DIAGRAM (ERD)
    # ==========================================================================
    add_heading1("6. Entity relationship diagram (ERD)")

    add_heading2("6.1 Conceptual data model and cardinalities")
    add_p("The FreshBasket data architecture uses 14 relational entities. The relationships enforce business integrity rules across the grocery ordering lifecycle. A user registers multiple saved delivery addresses in a 1 to M relationship between users and addresses. A category groups multiple produce items in a 1 to M relationship between categories and products. A product offers multiple measurement units such as kg, 500g, or dozen in a 1 to M relationship between products and product_units. A customer places multiple orders over time in a 1 to M relationship between users and orders. An order contains multiple line items linked to specific products through an M to N relationship resolved via the order_items junction table. An order generates multiple chronological status records in a 1 to M relationship between orders and order_status_history. Customers submit product reviews for delivered items in a 1 to M relationship between users/products and reviews. Customers file dispute tickets linked to orders in a 1 to M relationship between users/orders and complaints. Inventory adjustments write permanent records to an audit log in a 1 to M relationship between products and inventory_logs.")

    add_heading2("6.2 Entity relationship diagram")
    add_p("The diagram below illustrates the 14 data entities, primary keys, foreign key associations, and Crow's foot cardinalities.")
    add_image_box("diagrams/erd_diagram.png", "Entity Relationship Diagram (ERD) with 14 System Entities", width_in=6.2)

    add_heading2("6.3 Entity dictionary and attribute definitions")
    add_p("The table below summarizes the core entities, their primary keys, foreign key dependencies, and business roles.")

    ent_tbl = doc.add_table(rows=15, cols=4)
    ent_headers = ["Entity Name", "Primary Key", "Foreign Keys", "Description"]
    ent_data = [
        ["users", "id", "None", "Stores customer, admin, supervisor, and supplier profiles with hashed credentials."],
        ["addresses", "id", "user_id", "Stores delivery locations, street lines, GPS coordinates, and nearest hub IDs."],
        ["categories", "id", "None", "Organizes produce into Fruits, Vegetables, Leafy Greens, Organic, and Seasonal."],
        ["products", "id", "category_id", "Maintains produce catalog, freshness grade, harvest date, and base pricing."],
        ["product_units", "id", "product_id", "Configures purchasing units (kg, 500g, piece, dozen) with price multipliers."],
        ["inventory_logs", "id", "product_id, actor_id", "Records every stock increase or decrease with actor ID and audit reason."],
        ["orders", "id", "user_id, address_id", "Stores placed orders, total billing amounts, delivery slots, and current status."],
        ["order_items", "id", "order_id, product_id", "Stores line item records including selected unit, unit price, and quantity."],
        ["order_status_history", "id", "order_id, changed_by", "Maintains chronological records of 5-stage fulfillment status changes."],
        ["quality_checks", "id", "product_id, inspector_id", "Stores produce quality inspection logs, freshness scores, and lot numbers."],
        ["reviews", "id", "product_id, user_id, order_id", "Contains post-delivery star ratings (1 to 5) and customer feedback."],
        ["coupons", "id", "None", "Manages discount promotional vouchers with percentage or fixed deductions."],
        ["complaints", "id", "user_id, order_id", "Manages customer dispute tickets (FB-TKT-XXXXX) with SMTP email tracking."],
        ["audit_logs", "id", "user_id", "Records security and administrative system events for compliance auditing."]
    ]
    format_table(ent_tbl, [Inches(1.5), Inches(0.9), Inches(1.4), Inches(2.4)], ent_headers, ent_data)

    doc.add_page_break()

    # ==========================================================================
    # SECTION 7: DATABASE DESIGN
    # ==========================================================================
    add_heading1("7. Database design")

    add_heading2("7.1 Relational database schema structure")
    add_p("The FreshBasket database is designed in Third Normal Form (3NF). Every non-key attribute depends solely on the primary key. This structure eliminates data redundancies and update anomalies. Foreign keys enforce referential integrity with ON DELETE RESTRICT constraints on core transactional records (orders and products) to prevent accidental data loss.")
    add_p("The database uses SQLite via the native node:sqlite driver. SQLite runs within the server process. It requires zero background service configuration. Query execution is fast, with read latency under 5 milliseconds. The schema is portable and directly upgradeable to PostgreSQL or MySQL for large-scale enterprise deployments.")

    add_heading2("7.2 Physical relational schema diagram")
    add_p("The diagram below displays the normalized relational database tables, columns, data types, and primary-foreign key relationships.")
    add_image_box("diagrams/database_schema.png", "Physical Relational Database Schema Diagram (14 Tables)", width_in=6.2)

    add_heading2("7.3 Complete database table specifications")
    add_p("The data dictionaries below define the schema specifications for the primary database tables.")

    add_heading3("Table 1: users")
    t1 = doc.add_table(rows=8, cols=5)
    t1_headers = ["Column Name", "Data Type", "Null / Key", "Default", "Description"]
    t1_data = [
        ["id", "INTEGER", "PK, NOT NULL", "AUTOINCREMENT", "Unique internal user identifier"],
        ["full_name", "TEXT", "NOT NULL", "None", "User full legal name"],
        ["email", "TEXT", "UNIQUE, NOT NULL", "None", "Account email address used for login"],
        ["password_hash", "TEXT", "NOT NULL", "None", "Argon2 / bcrypt password hash string"],
        ["phone_number", "TEXT", "NULL", "None", "Contact phone number for delivery"],
        ["role", "TEXT", "NOT NULL", "'customer'", "Role: customer, admin, logistics, supplier"],
        ["created_at", "DATETIME", "NOT NULL", "CURRENT_TIMESTAMP", "Account registration timestamp"]
    ]
    format_table(t1, [Inches(1.3), Inches(0.9), Inches(1.1), Inches(1.2), Inches(1.7)], t1_headers, t1_data)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_heading3("Table 2: products")
    t2 = doc.add_table(rows=11, cols=5)
    t2_headers = ["Column Name", "Data Type", "Null / Key", "Default", "Description"]
    t2_data = [
        ["id", "INTEGER", "PK, NOT NULL", "AUTOINCREMENT", "Unique produce product identifier"],
        ["category_id", "INTEGER", "FK, NOT NULL", "None", "Foreign key references categories(id)"],
        ["name", "TEXT", "NOT NULL", "None", "Produce title (e.g., Farm Fresh Tomatoes)"],
        ["slug", "TEXT", "UNIQUE, NOT NULL", "None", "URL slug for routing"],
        ["base_unit", "TEXT", "NOT NULL", "'kg'", "Base measurement unit (kg, piece, dozen)"],
        ["base_price", "REAL", "NOT NULL", "0.0", "Base price per unit in PKR"],
        ["stock_quantity", "REAL", "NOT NULL", "0.0", "Current available inventory balance"],
        ["freshness_grade", "TEXT", "NOT NULL", "'A'", "Produce quality grade: A, B, C"],
        ["freshness_score", "INTEGER", "NOT NULL", "95", "Freshness quality index (1 to 100)"],
        ["farm_origin", "TEXT", "NULL", "None", "Harvest region and grower collective"]
    ]
    format_table(t2, [Inches(1.3), Inches(0.9), Inches(1.1), Inches(1.2), Inches(1.7)], t2_headers, t2_data)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_heading3("Table 3: orders")
    t3 = doc.add_table(rows=11, cols=5)
    t3_headers = ["Column Name", "Data Type", "Null / Key", "Default", "Description"]
    t3_data = [
        ["id", "INTEGER", "PK, NOT NULL", "AUTOINCREMENT", "Unique order identifier"],
        ["user_id", "INTEGER", "FK, NOT NULL", "None", "Foreign key references users(id)"],
        ["address_id", "INTEGER", "FK, NOT NULL", "None", "Foreign key references addresses(id)"],
        ["order_number", "TEXT", "UNIQUE, NOT NULL", "None", "Customer order code (FB-ORD-XXXXX)"],
        ["subtotal", "REAL", "NOT NULL", "0.0", "Sum of line item subtotals"],
        ["delivery_fee", "REAL", "NOT NULL", "0.0", "Doorstep dispatch fee"],
        ["total_amount", "REAL", "NOT NULL", "0.0", "Grand total payable upon delivery"],
        ["payment_type", "TEXT", "NOT NULL", "'cod'", "Payment method: Cash on Delivery"],
        ["order_status", "TEXT", "NOT NULL", "'pending'", "Stage: pending, confirmed, packed, out_for_delivery, delivered"],
        ["assigned_courier", "TEXT", "NULL", "None", "Name and contact number of courier rider"]
    ]
    format_table(t3, [Inches(1.3), Inches(0.9), Inches(1.1), Inches(1.2), Inches(1.7)], t3_headers, t3_data)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_heading3("Table 4: complaints")
    t4 = doc.add_table(rows=9, cols=5)
    t4_headers = ["Column Name", "Data Type", "Null / Key", "Default", "Description"]
    t4_data = [
        ["id", "INTEGER", "PK, NOT NULL", "AUTOINCREMENT", "Unique complaint dispute identifier"],
        ["ticket_number", "TEXT", "UNIQUE, NOT NULL", "None", "Ticket reference code (FB-TKT-XXXXX)"],
        ["user_id", "INTEGER", "FK, NULL", "None", "Foreign key references users(id)"],
        ["order_id", "INTEGER", "FK, NULL", "None", "Foreign key references orders(id)"],
        ["category", "TEXT", "NOT NULL", "'General'", "Category: Quality, Missing Item, Delay, Billing"],
        ["status", "TEXT", "NOT NULL", "'open'", "Ticket status: open, in_progress, resolved"],
        ["admin_reply", "TEXT", "NULL", "None", "Official response dispatched via Gmail SMTP"],
        ["created_at", "DATETIME", "NOT NULL", "CURRENT_TIMESTAMP", "Ticket creation timestamp"]
    ]
    format_table(t4, [Inches(1.3), Inches(0.9), Inches(1.1), Inches(1.2), Inches(1.7)], t4_headers, t4_data)

    doc.add_page_break()

    # ==========================================================================
    # SECTION 8: INTERFACE DESIGN
    # ==========================================================================
    add_heading1("8. Interface design")

    add_heading2("8.1 User interface layout and responsive principles")
    add_p("The FreshBasket interface uses modern responsive web principles. Built with React 18 and Tailwind CSS, the application provides an intuitive shopping workflow across mobile phones, tablets, and desktop workstations. The visual hierarchy uses clear produce images, price per unit indicators, farm origin tags, and freshness grades. The navigation adapts to screen dimensions. Mobile devices present a slide-over navigation drawer, accessible search fields, and touch-friendly cart triggers. The header displays the detected customer area and nearest cold-chain hub. Loading skeletons, toast alerts, and status badges confirm user actions.")

    add_heading2("8.2 Storefront catalog and hub selector")
    add_p("The storefront provides an interactive produce marketplace. Customers search produce by name, filter items using category tabs (All, Fruits, Vegetables, Leafy Greens, Organic, Seasonal), and select their regional hub. Produce cards present freshness badges and instant add-to-cart buttons.")
    add_image_box("screenshots/gui_1_storefront.png", "Storefront Catalog with Category Filters and Produce Cards", width_in=6.0)

    add_heading2("8.3 Produce details and farm origin display")
    add_p("Clicking a produce card opens the product detail view. The screen displays high-resolution imagery, freshness quality scores, harvest date, regional farm origin, nutritional metrics, and verified customer reviews.")
    add_image_box("screenshots/gui_2_product_detail.png", "Product Detail Screen Showing Nutritional Info and Farm Origin", width_in=6.0)

    add_heading2("8.4 Shopping cart and cash on delivery checkout")
    add_p("The checkout screen provides a multi-step purchase flow. Customers select measurement units, verify item pricing, pick delivery time slots, review the Geoapify doorstep map pin, and confirm their order with Cash on Delivery.")
    add_image_box("screenshots/gui_3_checkout_flow.png", "Shopping Cart and Cash on Delivery Checkout Screen", width_in=6.0)

    add_heading2("8.5 Real-time five-stage order tracking portal")
    add_p("The order tracking portal provides live order visibility. Customers view their order reference, delivery address, five-stage progress bar (Pending to Delivered), estimated delivery time, and route map showing dispatch courier transit.")
    add_image_box("screenshots/gui_4_order_tracker.png", "Real-Time Five-Stage Order Tracking Portal", width_in=6.0)

    add_heading2("8.6 Administrative dashboard and logistics console")
    add_p("The administrative and logistics console provides store management tools. Managers monitor gross revenue, order volume, low-stock alerts, customer complaints, review moderation, and dispatch courier assignments.")
    add_image_box("screenshots/gui_5_admin_dashboard.png", "Administrative Control Center and Logistics Dispatch Console", width_in=6.0)

    doc.save("task2.docx")
    doc.save("task1_documentation.docx")

    print(f"Generated task2.docx and task1_documentation.docx successfully.")
    if all_audit_errors:
        print(f"WARNING: Found {len(all_audit_errors)} audit errors:")
        for err in all_audit_errors:
            print("  -", err)
    else:
        print("✓ Zero compliance audit errors found. All strict writing rules passed!")

if __name__ == '__main__':
    build_task2_document()
